from ..config import UPLOADS
def inspect_file(path):
 p=(UPLOADS/path).resolve()
 if UPLOADS not in p.parents or not p.is_file():return {'error':'File not found'}
 ext=p.suffix.lower()
 try:
  if ext=='.pdf':
   from pypdf import PdfReader;r=PdfReader(str(p));return {'type':'pdf','pages':len(r.pages),'text':'\n'.join((x.extract_text() or '') for x in r.pages)[:50000]}
  if ext in ('.xlsx','.xlsm'):
   from openpyxl import load_workbook;wb=load_workbook(p,read_only=True,data_only=True);return {'type':'excel','sheets':{ws.title:[list(row) for row in list(ws.iter_rows(values_only=True))[:500] ] for ws in wb.worksheets}}
  if ext=='.csv':
   import pandas as pd;df=pd.read_csv(p);return {'type':'csv','rows':len(df),'columns':list(df.columns),'preview':df.head(50).to_dict(orient='records')}
  return {'type':'text','text':p.read_text(encoding='utf-8')[:50000]}
 except Exception as e:return {'error':str(e)}
